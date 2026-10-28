"""Cleaning report: before/after counts and Excel workbook export."""

from pathlib import Path

import pandas as pd

from src.utils.logger import get_logger

logger = get_logger(__name__)

CLEANED_TABLE_ORDER = ["trains", "stations", "routes", "delay_records", "delay_remarks"]


def build_cleaning_report(before_counts: dict, after_counts: dict, notes=None) -> dict:
    """Build the cleaning report dict with before/after counts per table."""
    return {
        "before_counts": dict(before_counts),
        "after_counts": dict(after_counts),
        "tables": {
            t: {"before": before_counts.get(t, 0), "after": after_counts.get(t, 0)}
            for t in CLEANED_TABLE_ORDER
        },
        "notes": list(notes or []),
    }


def _excel_rows(df: pd.DataFrame) -> list:
    """Rows of native Python values (None for missing), converted column-wise for speed."""
    cols = []
    for c in df.columns:
        col = df[c]
        if pd.api.types.is_datetime64_any_dtype(col):
            vals = [None if pd.isna(v) else v.to_pydatetime() for v in col]
        else:
            vals = col.astype(object).where(col.notna(), None).tolist()
        cols.append(vals)
    return [list(r) for r in zip(*cols)]


def write_cleaned_excel(data: dict, report: dict, path) -> None:
    """Write one Excel workbook: cleaned tables (one sheet each, with headers) + Cleaning_Report sheet."""
    from openpyxl import Workbook

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook(write_only=True)

    for table in CLEANED_TABLE_ORDER:
        if table not in data or data[table] is None:
            continue
        ws = wb.create_sheet(title=table[:31])
        df = data[table]
        ws.append([str(c) for c in df.columns])
        for row in _excel_rows(df):
            ws.append(row)
        logger.info("Excel sheet %s rows=%d", table, len(df))

    ws = wb.create_sheet(title="Cleaning_Report")
    ws.append(["table", "before_count", "after_count"])
    for table in CLEANED_TABLE_ORDER:
        ws.append([table, report.get("before_counts", {}).get(table, 0), report.get("after_counts", {}).get(table, 0)])

    wb.save(path)
    logger.info("Cleaning report Excel written to %s", path)
